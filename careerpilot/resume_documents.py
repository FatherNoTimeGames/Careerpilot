"""Local-only document extraction/rendering. Invoked in a resource-limited child process."""
import base64
import io
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

MAX_UPLOAD = 5 * 1024 * 1024
MAX_TEXT = 60000
MAX_PAGES = 20


def extract(raw, kind):
    if not raw or len(raw) > MAX_UPLOAD:
        raise ValueError('Upload must be between 1 byte and 5 MB')
    if kind == 'pdf':
        from pypdf import PdfReader
        if not raw.startswith(b'%PDF-'):
            raise ValueError('File does not contain a PDF')
        reader = PdfReader(io.BytesIO(raw), strict=True)
        if reader.is_encrypted:
            raise ValueError('Password-protected PDFs are not supported; upload a decrypted copy')
        if len(reader.pages) > MAX_PAGES:
            raise ValueError('Maximum 20 pages per resume')
        text = '\n\n'.join(page.extract_text() or '' for page in reader.pages)
    elif kind == 'docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            files = archive.infolist()
            if len(files) > 1000 or sum(f.file_size for f in files) > 25 * 1024 * 1024:
                raise ValueError('DOCX expands beyond allowed limits')
            if 'word/document.xml' not in archive.namelist():
                raise ValueError('File is not a DOCX document')
            if any('vbaproject' in f.filename.lower() for f in files):
                raise ValueError('Macro-enabled documents are not supported')
        doc = Document(io.BytesIO(raw))
        # Iterate blocks in reading order, including table cells; never execute links/macros.
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        parts = []
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                parts.append(block.text)
            elif isinstance(block, Table):
                for row in block.rows:
                    parts.append(' | '.join(cell.text for cell in row.cells))
        text = '\n'.join(parts)
        for section in doc.sections:
            text += '\n' + '\n'.join(p.text for p in section.header.paragraphs + section.footer.paragraphs if p.text)
    else:
        raise ValueError('Only PDF and DOCX are supported')
    text = text.replace('\x00', '').strip()
    if len(text) > MAX_TEXT:
        raise ValueError('Extracted text exceeds 60,000 characters')
    return {'text': text, 'warnings': [] if text else [
        'No readable text found. This may be a scanned or image-only document. OCR is not available; add facts manually.']}


def normalize(text):
    import unicodedata
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', text))


def export(blocks):
    from pypdf import PdfReader
    from docx import Document
    from docx.shared import Inches, Pt
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    import pypdfium2 as pdfium
    from PIL import Image, ImageChops

    fonts = Path(reportlab.__file__).parent / 'fonts'
    for name, filename in [('Resume', 'Vera.ttf'), ('ResumeBold', 'VeraBd.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    font = pdfmetrics.getFont('Resume')
    expected = []
    pdf = io.BytesIO()
    doc = SimpleDocTemplate(pdf, pagesize=(612, 792), rightMargin=50, leftMargin=50,
                            topMargin=50, bottomMargin=50, title='Resume', author='')
    body = ParagraphStyle('Body', fontName='Resume', fontSize=10, leading=14,
                          spaceAfter=7, alignment=TA_LEFT, splitLongWords=True)
    heading = ParagraphStyle('Heading', parent=body, fontName='ResumeBold', fontSize=12,
                             spaceBefore=12, spaceAfter=8, keepWithNext=True)
    word = Document()
    word.core_properties.author = ''
    word.core_properties.last_modified_by = ''
    word.sections[0].top_margin = word.sections[0].bottom_margin = Inches(.7)
    word.sections[0].left_margin = word.sections[0].right_margin = Inches(.7)
    word.styles['Normal'].font.name = 'Calibri'
    word.styles['Normal'].font.size = Pt(11)
    flow = []
    last_section = None
    for block in blocks:
        section, text = block['heading'], block['text']
        unsupported = {c for c in text + section if not c.isspace() and ord(c) not in font.face.charToGlyph}
        if unsupported:
            raise ValueError('Export font cannot represent some characters; revise unsupported symbols before export')
        if section != last_section:
            expected.append(section)
            flow.append(Paragraph(escape(section), heading))
            word.add_heading(section, level=1)
            last_section = section
        for line in text.splitlines():
            if line.strip():
                expected.append(line.strip())
                flow.append(Paragraph(escape(line.strip()), body))
                word.add_paragraph(line.strip())
    if not flow:
        raise ValueError('No confirmed content to export')
    doc.build(flow)
    raw_pdf = pdf.getvalue()
    reader = PdfReader(io.BytesIO(raw_pdf))
    if not 1 <= len(reader.pages) <= MAX_PAGES:
        raise ValueError('Export must contain 1–20 pages')
    extracted = '\n'.join(p.extract_text() or '' for p in reader.pages)
    if normalize(extracted) != normalize('\n'.join(expected)):
        raise ValueError('PDF completeness check failed; export withheld')
    renders = []
    with pdfium.PdfDocument(raw_pdf) as rendered:
        for page in rendered:
            try:
                bitmap = page.render(scale=1.25)
                try:
                    image = bitmap.to_pil().convert('RGB')
                    bounds = ImageChops.difference(image, Image.new('RGB', image.size, 'white')).getbbox()
                    if not bounds or bounds[0] < 40 or bounds[1] < 40 or bounds[2] > image.width-40 or bounds[3] > image.height-40:
                        raise ValueError('Rendered PDF margin/blank-page check failed; export withheld')
                    out = io.BytesIO()
                    image.save(out, format='PNG')
                    renders.append(base64.b64encode(out.getvalue()).decode())
                finally:
                    bitmap.close()
            finally:
                page.close()
    docx = io.BytesIO()
    word.save(docx)
    # Also round-trip editable content; Word itself is not available on this host.
    reopened = Document(io.BytesIO(docx.getvalue()))
    if normalize('\n'.join(p.text for p in reopened.paragraphs)) != normalize('\n'.join(expected)):
        raise ValueError('DOCX completeness check failed; export withheld')
    return {'pdf': base64.b64encode(raw_pdf).decode(), 'docx': base64.b64encode(docx.getvalue()).decode(),
            'pages': renders, 'checks': {'text_complete': True, 'docx_complete': True,
                                       'rendered_pages': len(renders), 'margins_checked': True,
                                       'human_visual_review_required': True}}


def main():
    # Bound untrusted document processing on supported Linux hosts. No network operations.
    import resource
    resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
    resource.setrlimit(resource.RLIMIT_AS, (1024**3, 1024**3))
    request = json.load(sys.stdin)
    try:
        result = extract(base64.b64decode(request['data'], validate=True), request['kind']) if request['action'] == 'extract' else export(request['blocks'])
        print(json.dumps({'ok': True, 'result': result}))
    except Exception as exc:
        # Do not send library exceptions containing document text to logs or clients.
        message = str(exc) if type(exc) is ValueError else 'Document processing failed; check format or add facts manually'
        print(json.dumps({'ok': False, 'error': message}))


if __name__ == '__main__':
    main()
