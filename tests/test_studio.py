"""Synthetic resume fixtures only; no candidate documents are stored in the repository."""
import base64
import io
import json
import os
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from unittest.mock import patch
from docx import Document
from reportlab.pdfgen import canvas
from pypdf import PdfReader
from careerpilot.store import connect
from careerpilot import studio
from careerpilot.resume_review import organize, profile_warnings
from careerpilot.security import initialize_password, owner_password
from careerpilot.web import server
from careerpilot.discovery import ingest
from test_pipeline import SOURCE, fixture


def docx_fixture():
    doc = Document()
    doc.add_paragraph('Example Candidate\nexample@example.invalid')
    doc.add_heading('Professional Experience', 1)
    doc.add_paragraph('IT Technician\nJuly 2025 – Present\nSupported Windows and Microsoft 365.')
    doc.add_heading('Education', 1)
    doc.add_paragraph('Bachelor of Arts\nExample College')
    doc.add_heading('Technical Skills', 1)
    doc.add_paragraph('Windows workstation support')
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def fact(text='Supported Windows and Microsoft 365.', status='confirmed', **extra):
    return {'id': 'skill-1', 'section': 'skills', 'text': text, 'status': status,
            'evidence': 'Synthetic test candidate explicitly confirmed this fact',
            'document_id': None, 'start_date': '', 'end_date': '', **extra}


class Studio(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dbpath = self.tmp.name+'/jobs.sqlite3'
        self.db = connect(self.dbpath)

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def profile(self, facts=None):
        revision = studio.save_profile(self.db, {'expected_revision': studio.current_profile(self.db)['revision'],
                                                'facts': facts or [fact()]})['revision']
        studio.verify_profile(self.db, {'expected_revision': revision, 'acknowledge': True})
        return revision

    def version(self, **extra):
        return studio.create_version(self.db, {'profile_revision': studio.current_profile(self.db)['revision'],
                                              'label': 'Test resume', **extra})['version_id']

    def preview_and_approve(self, ident):
        v = studio.get_version(self.db, ident)
        checks = studio.preview_version(self.db, {'version_id': ident, 'expected_edit_revision': v['edit_revision']})['checks']
        studio.approve_version(self.db, {'version_id': ident, 'expected_edit_revision': v['edit_revision'],
                                        'truth_reviewed': True, 'layout_reviewed': True})
        return checks

    def test_import_originals_and_uncertain_facts_never_auto_verify(self):
        raw = docx_fixture()
        result = studio.import_document(self.db, {'filename': 'resume.docx', 'data': base64.b64encode(raw).decode()})
        row = self.db.execute('SELECT * FROM resume_documents').fetchone()
        self.assertEqual(row['original'], raw)
        self.assertIn('Microsoft 365', row['extracted_text'])
        facts = result['suggested_facts']
        self.assertTrue(all(f['status'] == 'pending' for f in facts))
        experience = next(f for f in facts if f['section'] == 'employment')
        self.assertEqual(experience['start_date'], '2025-07')
        self.assertEqual(experience['end_date'], 'present')
        self.assertEqual(studio.current_profile(self.db)['revision'], 0)
        studio.save_profile(self.db, {'expected_revision': 0, 'facts': facts})
        with self.assertRaises(ValueError):
            self.version()
        with self.assertRaises(ValueError):
            studio.verify_profile(self.db, {'expected_revision': 1, 'acknowledge': True})
        # Both PDF extraction and DOCX extraction must organize readable source content.
        from careerpilot.resume_documents import export
        pdf = export([{'heading': 'Professional Experience', 'text': 'IT Technician\nSupported Windows workstations.'}])
        imported_pdf = studio.import_document(self.db, {'filename': 'resume.pdf', 'data': pdf['pdf']})
        self.assertEqual(imported_pdf['suggested_facts'][0]['section'], 'employment')
        self.assertEqual(self.db.execute('SELECT original FROM resume_documents WHERE id=?', (imported_pdf['document_id'],)).fetchone()[0], base64.b64decode(pdf['pdf']))
        self.assertEqual(studio.current_profile(self.db)['revision'], 1)

    def test_unreadable_scanned_and_bad_uploads_preserve_original(self):
        stream = io.BytesIO()
        pdf = canvas.Canvas(stream)
        pdf.rect(50, 50, 50, 50)
        pdf.showPage()
        pdf.save()
        for filename, raw in [('scan.pdf', stream.getvalue()), ('broken.docx', b'not a zip')]:
            result = studio.import_document(self.db, {'filename': filename, 'data': base64.b64encode(raw).decode()})
            self.assertTrue(result['warnings'])
            self.assertFalse(result['suggested_facts'])
            self.assertEqual(self.db.execute('SELECT original FROM resume_documents WHERE id=?', (result['document_id'],)).fetchone()[0], raw)
        with self.assertRaises(ValueError):
            studio.import_document(self.db, {'filename': 'unsafe.exe', 'data': 'YQ=='})
        with self.assertRaises(ValueError):
            studio.import_document(self.db, {'filename': 'huge.pdf', 'data': base64.b64encode(b'a'*(5*1024*1024+1)).decode()})

    def test_date_contradictions_and_concurrency_are_blocked(self):
        bad = fact('IT Technician 2025 – 2022', section='employment', start_date='2025', end_date='2022')
        studio.save_profile(self.db, {'expected_revision': 0, 'facts': [bad]})
        self.assertTrue(profile_warnings([bad])['blockers'])
        with self.assertRaises(ValueError):
            studio.verify_profile(self.db, {'expected_revision': 1, 'acknowledge': True})
        with self.assertRaises(studio.Conflict):
            studio.save_profile(self.db, {'expected_revision': 0, 'facts': [fact()]})
        self.assertTrue(profile_warnings([fact('July 2025 – Present',start_date='2024-01')])['blockers'])

    def test_confirmed_only_versions_and_manual_number_credential_guard(self):
        self.profile([fact(), fact('Personal hobby', status='inferred', id='inferred', section='personal'),
                      fact('Unknown accomplishment', status='pending', id='pending')])
        ident = self.version()
        v = studio.get_version(self.db, ident)
        self.assertEqual(len(json.loads(v['blocks'])), 1)
        with self.assertRaises(ValueError):
            self.version(fact_ids=['pending'])
        result = studio.edit_version(self.db, {'version_id': ident, 'expected_edit_revision': 1,
             'blocks': [{'fact_id':'skill-1', 'revised':'Certified CCNA; reduced incidents by 95%.', 'explanation':'Invented test claim'}]})
        self.assertGreaterEqual(len(result['review']['blockers']), 2)
        with self.assertRaises(ValueError):
            studio.preview_version(self.db, {'version_id': ident, 'expected_edit_revision': 2})
        with self.assertRaises(studio.Conflict):
            studio.edit_version(self.db, {'version_id': ident, 'expected_edit_revision': 1, 'blocks': []})

    def test_approval_master_job_selection_and_immutable_snapshots(self):
        revision = self.profile()
        ident = self.version()
        with self.assertRaises(ValueError):
            studio.approve_version(self.db, {'version_id': ident, 'expected_edit_revision': 1,
                                            'truth_reviewed':True, 'layout_reviewed':True})
        checks = self.preview_and_approve(ident)
        self.assertTrue(checks['text_complete'])
        self.assertIsNone(studio.state(self.db)['master_version_id'])
        original_pdf = self.db.execute('SELECT pdf FROM resume_exports WHERE version_id=?', (ident,)).fetchone()[0]
        with self.assertRaises(ValueError):
            studio.set_master(self.db, {'version_id': ident, 'expected_master_version_id': None})
        studio.set_master(self.db, {'version_id': ident, 'expected_master_version_id':None, 'approve_replacement':True})
        ingest(self.db, SOURCE, [fixture()])
        studio.select_for_job(self.db, {'version_id': ident, 'job_id': 1})
        self.assertEqual(self.db.execute('SELECT review_status FROM jobs WHERE id=1').fetchone()[0], 'discovered')
        self.assertEqual(studio.state(self.db)['job_resumes'][0]['version_id'], ident)
        self.profile([fact('Added verified Active Directory support.', id='new')])
        self.assertEqual(self.db.execute('SELECT pdf FROM resume_exports WHERE version_id=?', (ident,)).fetchone()[0], original_pdf)
        self.assertEqual(studio.get_version(self.db, ident)['profile_revision'], revision)
        self.assertTrue(studio.state(self.db)['versions'][0]['profile_outdated'])
        with self.assertRaises(ValueError):
            studio.edit_version(self.db, {'version_id':ident,'expected_edit_revision':1,'blocks':[]})
        restarted = connect(self.dbpath)
        try:
            state = studio.state(restarted)
            self.assertEqual(state['master_version_id'], ident)
            self.assertEqual(state['profile']['facts'][0]['text'], 'Added verified Active Directory support.')
        finally:
            restarted.close()

    def test_stale_preview_invalidated_and_outdated_profile_draft_blocked(self):
        self.profile()
        ident = self.version()
        studio.preview_version(self.db, {'version_id':ident,'expected_edit_revision':1})
        studio.edit_version(self.db, {'version_id':ident,'expected_edit_revision':1,'blocks':[
            {'fact_id':'skill-1','revised':'Provided Windows and Microsoft 365 support.','explanation':'Use clearer wording without new responsibilities.'}]})
        self.assertIsNone(studio.state(self.db)['versions'][0]['preview'])
        self.assertEqual(self.db.execute('SELECT count(*) FROM resume_pages').fetchone()[0], 0)
        studio.preview_version(self.db, {'version_id':ident,'expected_edit_revision':2})
        self.profile([fact('Different verified facts')])
        with self.assertRaises(ValueError):
            studio.approve_version(self.db, {'version_id':ident,'expected_edit_revision':2,'truth_reviewed':True,'layout_reviewed':True})

    def test_multi_page_pdf_render_text_and_docx_roundtrip(self):
        facts = [fact(('Supported employee Windows workstations and documented troubleshooting steps. ' * 8).strip(),
                      id='experience-'+str(i), section='employment') for i in range(14)]
        self.profile(facts)
        ident = self.version()
        checks = self.preview_and_approve(ident)
        self.assertGreater(checks['rendered_pages'], 1)
        exported = self.db.execute('SELECT pdf,docx FROM resume_exports WHERE version_id=?',(ident,)).fetchone()
        pdf = PdfReader(io.BytesIO(exported['pdf']))
        self.assertEqual(len(pdf.pages), checks['rendered_pages'])
        self.assertEqual(''.join(''.join(p.extract_text() or '' for p in pdf.pages).split()).count('documentedtroubleshootingsteps.'), 14*8)
        self.assertEqual(sum(p.text.count('documented troubleshooting steps.') for p in Document(io.BytesIO(exported['docx'])).paragraphs), 14*8)
        self.assertEqual(self.db.execute('SELECT count(*) FROM resume_pages').fetchone()[0], checks['rendered_pages'])

    def test_local_password_bootstrap_is_private_and_repeatable(self):
        with patch.dict(os.environ, {'CAREERPILOT_PASSWORD':''}):
            path = initialize_password(self.dbpath)
            first = owner_password(self.dbpath)
            self.assertGreaterEqual(len(first), 32)
            self.assertEqual(path.stat().st_mode & 0o777, 0o600)
            initialize_password(self.dbpath)
            self.assertEqual(owner_password(self.dbpath), first)


class StudioHTTP(unittest.TestCase):
    def test_private_routes_upload_download_and_no_leak_in_public_state(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, {'CAREERPILOT_PASSWORD':'synthetic-local-test-password'}):
            app = server(tmp+'/db', port=0)
            thread=threading.Thread(target=app.serve_forever)
            thread.start()
            base='http://127.0.0.1:'+str(app.server_port)
            auth='Basic '+base64.b64encode(b'careerpilot:synthetic-local-test-password').decode()
            opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
            def request(path, body=None, authorized=True):
                headers={'Authorization':auth} if authorized else {}
                if body is not None:
                    headers.update({'Content-Type':'application/json','X-CareerPilot':'1'})
                req=urllib.request.Request(base+path,data=json.dumps(body).encode() if body is not None else None,headers=headers)
                return opener.open(req)
            try:
                raw=docx_fixture()
                imported=json.load(request('/api/studio/import',{'filename':'resume.docx','data':base64.b64encode(raw).decode()}))
                ident=imported['document_id']
                self.assertEqual(request(f'/api/studio/documents/{ident}/original').read(),raw)
                with self.assertRaises(urllib.error.HTTPError) as error:
                    request(f'/api/studio/documents/{ident}/original',authorized=False)
                self.assertEqual(error.exception.code,401)
                state=json.load(request('/api/state'))
                self.assertNotIn('documents',state)
                self.assertNotIn('Example Candidate',json.dumps(state))
                profile=json.load(request('/api/studio/profile/save',{'expected_revision':0,'facts':[fact()]}))
                request('/api/studio/profile/verify',{'expected_revision':profile['revision'],'acknowledge':True}).close()
                ident=json.load(request('/api/studio/version/create',{'profile_revision':profile['revision'],'label':'Private test'}))['version_id']
                with self.assertRaises(urllib.error.HTTPError):
                    request(f'/api/studio/versions/{ident}/pdf')
                request('/api/studio/version/preview',{'version_id':ident,'expected_edit_revision':1}).close()
                self.assertTrue(request(f'/api/studio/versions/{ident}/pages/1').read().startswith(b'\x89PNG'))
                request('/api/studio/version/approve',{'version_id':ident,'expected_edit_revision':1,'truth_reviewed':True,'layout_reviewed':True}).close()
                self.assertTrue(request(f'/api/studio/versions/{ident}/pdf').read().startswith(b'%PDF-'))
                self.assertTrue(request(f'/api/studio/versions/{ident}/docx').read().startswith(b'PK'))
            finally:
                app.shutdown();thread.join();app.server_close()

    def test_studio_disabled_without_authentication_even_on_loopback(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ,{'CAREERPILOT_PASSWORD':''}):
            app=server(tmp+'/db',port=0)
            thread=threading.Thread(target=app.serve_forever);thread.start()
            try:
                opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
                with self.assertRaises(urllib.error.HTTPError) as error:
                    opener.open('http://127.0.0.1:'+str(app.server_port)+'/api/studio')
                self.assertEqual(error.exception.code,403)
            finally:
                app.shutdown();thread.join();app.server_close()
