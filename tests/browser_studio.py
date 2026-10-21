"""Optional real mobile browser check. Run from the checkout root with PYTHONPATH=."""
import base64
import io
import json
import os
import secrets
import sys
import tempfile
import threading
from pathlib import Path
from docx import Document
from playwright.sync_api import sync_playwright
from careerpilot.store import connect
from careerpilot.web import server
from careerpilot.discovery import ingest

# All browser data is synthetic, in a temporary DB, never the live candidate profile.
with tempfile.TemporaryDirectory() as tmp:
    os.environ['CAREERPILOT_PASSWORD']=secrets.token_urlsafe(32)
    dbpath=tmp+'/jobs.sqlite3'
    db=connect(dbpath)
    ingest(db,{'provider':'greenhouse','board':'synthetic','employer':'Test employer'},[{'external_id':'1','url':'https://example.invalid/jobs/1','employer':'Test employer','title':'IT Support Specialist','location':'Remote, United States','employment_type':'FullTime','is_remote':True,'description':'Synthetic internal support role.'}])
    db.close()
    httpd=server(dbpath,port=0)
    thread=threading.Thread(target=httpd.serve_forever);thread.start()
    try:
        with sync_playwright() as p:
            browser=p.chromium.launch(executable_path=os.environ.get('CAREERPILOT_CHROMIUM', '/usr/bin/chromium'),headless=True,args=['--no-sandbox'])
            context=browser.new_context(viewport={'width':393,'height':851},device_scale_factor=1,
                http_credentials={'username':'careerpilot','password':os.environ['CAREERPILOT_PASSWORD']},accept_downloads=True)
            page=context.new_page()
            errors=[]
            page.on('pageerror',lambda e:errors.append(str(e)))
            page.goto('http://127.0.0.1:'+str(httpd.server_port)+'/#studio')
            page.get_by_role('heading',name='Résumé Studio',exact=True).wait_for()
            doc=Document();doc.add_paragraph('Example Candidate\nexample@example.invalid')
            doc.add_heading('Technical Skills',1);doc.add_paragraph('Supported Windows workstations and Microsoft 365.')
            stream=io.BytesIO();doc.save(stream)
            page.locator('input[type=file]').set_input_files({'name':'synthetic-resume.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':stream.getvalue()})
            page.get_by_role('button',name='Import & extract').click()
            page.get_by_text('Original documents (1)',exact=True).wait_for()
            page.get_by_text('Original documents (1)',exact=True).click()
            page.get_by_role('button',name='Review extracted entries in profile').click()
            page.get_by_role('heading',name='Career profile',exact=True).wait_for()
            for fact in page.locator('[data-fact]').all():
                fact.locator('[name=evidence]').fill('Confirmed against synthetic original by test owner')
                fact.locator('[name=status]').select_option('confirmed')
            page.get_by_role('button',name='Save profile',exact=True).click()
            page.get_by_text('Revision 1 · Needs review',exact=True).wait_for()
            page.locator('#verify-profile input').check()
            page.get_by_role('button',name='Verify saved profile').click()
            page.get_by_text('Revision 1 · Verified',exact=True).wait_for()
            page.locator('nav a[href="#studio"]').click()
            page.locator('#create-version input[name=label]').fill('Synthetic browser résumé')
            page.get_by_role('button',name='Create review draft').click()
            page.locator('.version[open]').wait_for()
            skills=page.locator('.revision-block').filter(has_text='Technical Skills')
            skills.locator('textarea[name=revised]').fill('Provided Windows workstation and Microsoft 365 support.')
            skills.locator('textarea[name=explanation]').fill('Clearer wording; same verified support experience.')
            page.locator('nav a[href="#profile"]').click()
            page.get_by_role('heading',name='Career profile',exact=True).wait_for()
            page.locator('nav a[href="#studio"]').click()
            page.locator('.version[open]').wait_for()
            assert page.locator('.revision-block').filter(has_text='Technical Skills').locator('textarea[name=revised]').input_value()=='Provided Windows workstation and Microsoft 365 support.'
            page.get_by_role('button',name='Save wording & run checks').click()
            page.get_by_role('button',name='Check & render PDF preview').click()
            page.locator('.pdf-pages img').wait_for()
            page.wait_for_function("[...document.querySelectorAll('.pdf-pages img')].every(i => i.complete && i.naturalWidth > 0)")
            page.locator('.pdf-pages img').screenshot(path='/tmp/careerpilot-resume-render.png')
            page.screenshot(path='/tmp/careerpilot-studio-mobile.png',full_page=True)
            assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth'), 'Mobile horizontal overflow'
            page.locator('input[name=truth_reviewed]').check();page.locator('input[name=layout_reviewed]').check()
            page.get_by_role('button',name='Approve this version',exact=True).click()
            page.get_by_role('link',name='Download PDF',exact=True).wait_for()
            with page.expect_download() as download:
                page.get_by_role('link',name='Download PDF',exact=True).click()
            downloaded=Path(download.value.path()).read_bytes()
            assert downloaded.startswith(b'%PDF-')
            page.once('dialog',lambda dialog:dialog.accept())
            page.get_by_role('button',name='Make master résumé').click()
            page.get_by_role('button',name='Current master').wait_for()
            page.locator('.select-version select').select_option('1')
            page.get_by_role('button',name='Select for job',exact=True).click()
            page.get_by_text('IT Support Specialist · Test employer — Synthetic browser résumé (v1)',exact=True).wait_for()
            page.reload();page.locator('.version > summary').filter(has_text='MASTER').click();page.get_by_role('button',name='Current master').wait_for()
            assert not errors,errors
            print('Mobile browser flow passed: DOCX import, fact verification, manual revision, rendered preview, approval, PDF download, explicit master replacement and reload persistence. No JS errors or horizontal overflow.')
            browser.close()
    finally:
        httpd.shutdown();thread.join();httpd.server_close()
